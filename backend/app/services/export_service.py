import json
from app.models.transcript import TranscriptSegment, SessionResult

def format_seconds_to_srt(seconds: float) -> str:
    """
    초 단위 실수를 SRT 시간 포맷(HH:MM:SS,mmm)으로 변환합니다.
    """
    hrs = int(seconds // 3600)
    mins = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    millis = int(round((seconds - int(seconds)) * 1000))
    
    # 예외 방지 밀리초 자릿수 맞춤
    if millis >= 1000:
        millis = 999
        
    return f"{hrs:02d}:{mins:02d}:{secs:02d},{millis:03d}"

def export_to_srt(segments: list[TranscriptSegment]) -> str:
    """
    세그먼트 리스트를 SRT 자막 파일 형식 문자열로 변환합니다.
    """
    lines = []
    for i, seg in enumerate(segments, 1):
        start_str = format_seconds_to_srt(seg.start)
        end_str = format_seconds_to_srt(seg.end)
        lines.append(f"{i}")
        lines.append(f"{start_str} --> {end_str}")
        lines.append(f"[{seg.speaker}] {seg.text}")
        lines.append("")  # 공백 라인 구분
    return "\n".join(lines)

def export_to_txt(segments: list[TranscriptSegment]) -> str:
    """
    세그먼트 리스트를 텍스트 회의록 형식 문자열로 변환합니다.
    """
    lines = []
    for seg in segments:
        start_str = format_seconds_to_srt(seg.start).split(",")[0]  # 밀리초 절삭
        lines.append(f"[{start_str}] {seg.speaker}: {seg.text}")
    return "\n".join(lines)

def export_to_json(session: SessionResult) -> str:
    """
    세션 전체 결과를 JSON 포맷 문자열로 직렬화합니다.
    """
    return json.dumps({
        "metadata": {
            "session_id": session.session_id,
            "status": session.status,
            "created_at": session.created_at.isoformat() if session.created_at else None,
            "duration_sec": session.duration_sec,
            "speaker_count": session.speaker_count,
            "audio_language": session.audio_language,
        },
        "segments": [
            {
                "id": seg.id,
                "timestamp": {
                    "start": seg.start,
                    "end": seg.end,
                    "start_human": format_seconds_to_srt(seg.start),
                    "end_human": format_seconds_to_srt(seg.end),
                },
                "speaker": seg.speaker,
                "text": seg.text,
                "confidence": seg.confidence,
                "is_final": seg.is_final,
            }
            for seg in session.segments
        ]
    }, indent=2, ensure_ascii=False)


def export_to_docx(session: SessionResult) -> bytes:
    """
    세션 데이터를 DOCX 형식으로 생성합니다.
    requires: python-docx
    """
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise ImportError("python-docx is required for DOCX export. Install with: pip install python-docx")

    doc = Document()

    # 제목
    title = doc.add_heading('회의 기록', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 메타데이터
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.style = 'Light Grid Accent 1'
    cells = meta_table.rows[0].cells
    cells[0].text = '세션 ID'
    cells[1].text = session.session_id

    cells = meta_table.rows[1].cells
    cells[0].text = '화자 수'
    cells[1].text = str(session.speaker_count)

    cells = meta_table.rows[2].cells
    cells[0].text = '총 시간'
    cells[1].text = f"{session.duration_sec:.1f}초"

    cells = meta_table.rows[3].cells
    cells[0].text = '작성 시간'
    cells[1].text = session.created_at.isoformat() if session.created_at else 'Unknown'

    cells = meta_table.rows[4].cells
    cells[0].text = '상태'
    cells[1].text = session.status

    # 세그먼트 섹션
    doc.add_heading('발언 기록', level=1)

    for seg in session.segments:
        # 헤더: 화자, 시간, 신뢰도
        header = doc.add_paragraph()
        header_run = header.add_run(f"{seg.speaker}  [{format_seconds_to_srt(seg.start)} ~ {format_seconds_to_srt(seg.end)}]  신뢰도: {seg.confidence*100:.0f}%")
        header_run.bold = True
        header_run.font.size = Pt(11)

        # 텍스트 본문
        text_para = doc.add_paragraph(seg.text)
        text_para.paragraph_format.left_indent = Inches(0.25)

        # 최종 여부 표시
        if seg.is_final:
            status_run = text_para.add_run("  ✓")
            status_run.font.color.rgb = RGBColor(34, 197, 94)  # green

    # DOCX를 바이트로 반환
    from io import BytesIO
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
